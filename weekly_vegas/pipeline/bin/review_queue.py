"""Spreadsheet review queue for post drafts.

The user is spreadsheet-driven: export post_queue drafts to xlsx, mark
Status (approved/rejected) + an optional note, then apply the decisions
back to the DB. Only rows exported as drafts are ever touched, and apply
refuses to clobber a row whose DB status moved on since export.

Usage:
  python3 bin/review_queue.py export [--out PATH]
  python3 bin/review_queue.py apply --in PATH

Export columns: ID (hidden) | Type | Title | Text | Status | Note
  Status has a dropdown: draft / approved / rejected.
"""
import argparse
import os
import sys

BASE = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))  # weekly_vegas/pipeline root
sys.path.insert(0, BASE)
sys.path.insert(0, os.path.expanduser("~/workspace/skills/supabase-football-signal/bin"))
import sbclient  # noqa: E402

STATUSES = ["draft", "approved", "rejected"]
DEFAULT_OUT = os.path.join(BASE, "data", "review_queue.xlsx")


def _drafts():
    rows = sbclient.get_all(
        "post_queue",
        "?status=eq.draft&select=id,post_type,post_text,note,created_at"
        "&order=created_at.asc&limit=200")
    return rows


def _title(r):
    if r["post_type"] == "thread":
        n = r["post_text"].count("\n---\n") + 1
        return f"THREAD ({n} tweets)"
    first = (r["post_text"] or "").split("\n")
    return " / ".join(first[:2])[:60]


def export(out=DEFAULT_OUT):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment
    from openpyxl.worksheet.datavalidation import DataValidation
    from openpyxl.utils import get_column_letter

    rows = _drafts()
    wb = Workbook()
    ws = wb.active
    ws.title = "Review"
    headers = ["ID", "Type", "Title", "Text", "Status", "Reviewer note"]
    ws.append(headers)
    for r in rows:
        text = (r["post_text"] or "").replace("\n---\n", "\n\n")
        ws.append([r["id"], r["post_type"], _title(r), text, "draft",
                   r.get("note") or ""])
    # formatting
    ws.column_dimensions["A"].hidden = True
    ws.column_dimensions["B"].width = 10
    ws.column_dimensions["C"].width = 40
    ws.column_dimensions["D"].width = 60
    ws.column_dimensions["E"].width = 12
    ws.column_dimensions["F"].width = 40
    for row in ws.iter_rows(min_row=2, max_row=ws.max_row, max_col=6):
        row[3].alignment = Alignment(wrap_text=True, vertical="top")
        ws.row_dimensions[row[0].row].height = 60
    dv = DataValidation(type="list", formula1='"draft,approved,rejected"',
                        allow_blank=False)
    dv.error = "Pick draft, approved, or rejected"
    ws.add_data_validation(dv)
    dv.add(f"E2:E{ws.max_row}")
    # instructions sheet
    ins = wb.create_sheet("How to use")
    ins["A1"] = ("Mark each row's Status as approved or rejected and "
                 "optionally add a note, then run: "
                 "python3 bin/review_queue.py apply --in " + out)
    ins.column_dimensions["A"].width = 100
    _add_signals_sheet(wb)
    _add_scoreboard_sheet(wb)
    _add_quota_sheet(wb)
    wb.save(out)
    print(f"exported {len(rows)} drafts -> {out}")
    return out


def _add_signals_sheet(wb):
    """Read-only: all post-worthy signals with their numbers."""
    import json
    ws = wb.create_sheet("Signals")
    ws.append(["Player", "Pos", "Team", "Vegas PPR", "Expert PPR",
               "Delta", "Lean", "Rank gap"])
    path = os.path.join(BASE, "data", "signals_v4.json")
    try:
        sigs = json.load(open(path))
    except FileNotFoundError:
        sigs = []
    pw = [s for s in sigs if s.get("post_worthy")]
    pw.sort(key=lambda s: s.get("delta_ppr", 0))
    for s in pw:
        ws.append([s.get("player"), s.get("pos"), s.get("team"),
                   round(s.get("vegas_ppr", 0), 1),
                   round(s.get("expert_ppr", 0), 1),
                   round(s.get("delta_ppr", 0), 1),
                   "fade" if s.get("delta_ppr", 0) < 0 else "love",
                   s.get("rank_gap")])
    for col, w in zip("ABCDEFGH", [22, 6, 6, 11, 11, 8, 8, 10]):
        ws.column_dimensions[col].width = w
    ws.freeze_panes = "A2"


def _add_scoreboard_sheet(wb):
    """Read-only: calibration state from ranker_performance (live)."""
    ws = wb.create_sheet("Scoreboard")
    ws.append(["Source", "Format", "Season", "Week", "n", "MAE", "RMSE",
               "Bias"])
    try:
        rows = mgmt_query(
            """SELECT r.display_name AS source,
                      rp.metadata->>'scoring_format' AS fmt,
                      rp.season, rp.week,
                      (rp.metadata->>'n')::int AS n,
                      rp.accuracy_score AS mae,
                      (rp.metadata->>'rmse')::numeric AS rmse,
                      (rp.metadata->>'bias')::numeric AS bias
               FROM ranker_performance rp
               JOIN rankers r ON r.id = rp.ranker_id
               ORDER BY rp.season DESC, rp.week DESC, source;""")
        for x in rows:
            ws.append([x["source"], x["fmt"], x["season"], x["week"],
                       x["n"], float(x["mae"]), float(x["rmse"]),
                       float(x["bias"])])
        if not rows:
            ws.append(["no calibration rows yet — run bin/calibrate.py "
                       "after games go final"])
    except Exception as e:
        ws.append([f"query failed: {e}"])
    for col, w in zip("ABCDEFGH", [16, 10, 8, 6, 6, 8, 8, 8]):
        ws.column_dimensions[col].width = w
    ws.freeze_panes = "A2"


def mgmt_query(sql):
    sys.path.insert(0, os.path.expanduser(
        "~/workspace/skills/supabase-mgmt/bin"))
    import mgmt
    return mgmt.query(sql)


def _add_quota_sheet(wb):
    """Read-only: odds quota + pipeline health."""
    import json
    ws = wb.create_sheet("Quota & Health")
    ws.append(["Metric", "Value"])
    qpath = os.path.join(BASE, "data", "odds_cache", "quota.json")
    try:
        q = json.load(open(qpath))
        ws.append(["Odds API remaining", q.get("remaining")])
        ws.append(["Odds API used", q.get("used")])
        ws.append(["Quota last updated", q.get("at")])
    except FileNotFoundError:
        ws.append(["Odds API quota", "no quota file yet"])
    try:
        n = sbclient.count("odds_history")
        ws.append(["odds_history rows", n])
    except Exception as e:
        ws.append(["odds_history rows", f"query failed: {e}"])
    try:
        n = sbclient.count("projection_snapshots")
        ws.append(["projection_snapshots rows", n])
    except Exception as e:
        ws.append(["projection_snapshots rows", f"query failed: {e}"])
    ws.append(["Free-tier reserve rule", "keep 60 credits; full pull ~90+"])
    ws.column_dimensions["A"].width = 26
    ws.column_dimensions["B"].width = 40


def apply(path):
    from openpyxl import load_workbook
    wb = load_workbook(path)
    ws = wb["Review"]
    changed = skipped = 0
    for row in ws.iter_rows(min_row=2, values_only=True):
        rid, status, note = row[0], (row[4] or "draft").strip(), row[5] or ""
        if status == "draft" or not rid:
            continue
        if status not in ("approved", "rejected"):
            print(f"skip {rid[:8]}: bad status '{status}'")
            skipped += 1
            continue
        cur = sbclient.get("post_queue", f"?id=eq.{rid}&select=status")
        if not cur:
            print(f"skip {rid[:8]}: row gone")
            skipped += 1
            continue
        if cur[0]["status"] != "draft":
            print(f"skip {rid[:8]}: DB status is now {cur[0]['status']} (not clobbering)")
            skipped += 1
            continue
        patch = {"status": status}
        if note:
            patch["note"] = note
        sbclient.patch("post_queue", patch, f"?id=eq.{rid}")
        print(f"{status}: {rid[:8]}")
        changed += 1
    print(f"apply done: {changed} updated, {skipped} skipped")


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("cmd", choices=["export", "apply"])
    ap.add_argument("--out", default=DEFAULT_OUT)
    ap.add_argument("--in", dest="inp", default=DEFAULT_OUT)
    a = ap.parse_args()
    if a.cmd == "export":
        export(a.out)
    else:
        apply(a.inp)


if __name__ == "__main__":
    main()
